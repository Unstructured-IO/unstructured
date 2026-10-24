"""Literal missing-value tokens survive Excel parsing; blank and error cells do not."""

import base64
import io
import zlib
from datetime import datetime
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
import pandas as pd
import pytest

from unstructured.documents.elements import Table
from unstructured.partition.auto import partition
from unstructured.partition.xlsx import open_excel_workbook, partition_xlsx


@pytest.fixture
def token_workbook():
    book = openpyxl.Workbook()
    for row in (
        ["Item", "NA"],
        ["Pear", "N/A"],
        ["Fig", "null"],
        [None, None],
        ["Plum", "None"],
        ["Kiwi", "NaN"],
    ):
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("find_subtable", [False, True])
@pytest.mark.parametrize("include_header", [False, True])
@pytest.mark.parametrize("infer_table_structure", [False, True])
def test_literal_tokens_preserve_tables_and_blank_separators(
    token_workbook, find_subtable, include_header, infer_table_structure
):
    elements = partition_xlsx(
        file=io.BytesIO(token_workbook),
        find_subtable=find_subtable,
        include_header=include_header,
        infer_table_structure=infer_table_structure,
    )
    assert all(isinstance(element, Table) for element in elements)
    assert [element.text for element in elements] == (
        ["Item NA Pear N/A Fig null", ("Item NA " if include_header else "") + "Plum None Kiwi NaN"]
        if find_subtable
        else ["Item NA Pear N/A Fig null Plum None Kiwi NaN"]
    )
    if infer_table_structure:
        assert "N/A</" in elements[0].metadata.text_as_html
        assert "NaN</" in elements[-1].metadata.text_as_html
        if not find_subtable:
            assert "<tr><td/><td/></tr>" in elements[0].metadata.text_as_html
    else:
        assert all(element.metadata.text_as_html is None for element in elements)


def test_auto_dispatch_preserves_literal_tokens(token_workbook, tmp_path):
    path = tmp_path / "tokens.xlsx"
    path.write_bytes(token_workbook)
    assert [element.text for element in partition(filename=str(path))] == [
        "Item NA Pear N/A Fig null",
        "Plum None Kiwi NaN",
    ]


@pytest.mark.parametrize("compact_empty_rows", [False, True])
def test_reader_distinguishes_literal_tokens_blanks_errors_and_native_values(compact_empty_rows):
    tokens = ["NA", "N/A", "None", "null", "NaN", "nan", "NULL", "<NA>", "#N/A"]
    values = [*tokens, None, "", 0, False, datetime(2024, 2, 3), "end"]
    book = openpyxl.Workbook()
    for index, value in enumerate(values, 1):
        book.active.cell(index, 1, f"row {index}")
        cell = book.active.cell(index, 2, value)
        if value == "#N/A":
            cell.data_type = "s"
    book.active.cell(1, 3, "#N/A")  # A genuine error, unlike the text-typed token.
    stream = io.BytesIO()
    book.save(stream)
    with open_excel_workbook(stream.getvalue(), compact_empty_rows=compact_empty_rows) as reader:
        frame = reader.parse("Sheet")
    assert frame.iloc[: len(tokens), 1].tolist() == tokens
    assert pd.isna(frame.iloc[len(tokens), 1])
    assert pd.isna(frame.iloc[len(tokens) + 1, 1])
    assert frame.iloc[len(tokens) + 2, 1] == 0
    assert not frame.iloc[len(tokens) + 3, 1]  # pandas can infer False as numeric zero.
    assert not pd.isna(frame.iloc[len(tokens) + 3, 1])
    assert frame.iloc[len(tokens) + 4, 1] == datetime(2024, 2, 3)
    assert pd.isna(frame.iloc[0, 2])


def test_numeric_token_column_keeps_literal_and_integer_rendering():
    book = openpyxl.Workbook()
    for row in (["Item", "Count"], ["Pear", 1], ["Fig", "N/A"], ["Plum", 3]):
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    (table,) = partition_xlsx(file=io.BytesIO(stream.getvalue()), include_header=True)
    assert table.text == "Item Count Pear 1 Fig N/A Plum 3"
    assert "<td>1</td>" in table.metadata.text_as_html
    assert "<td>3</td>" in table.metadata.text_as_html


def test_reader_uses_cached_formula_results_without_recalculation():
    book = openpyxl.Workbook()
    for index, formula in enumerate(['="N/A"', "=7", "=NA()", '=""'], 1):
        book.active.cell(index, 1, f"row {index}")
        book.active.cell(index, 2, formula)
    stream = io.BytesIO()
    book.save(stream)
    # Supply known cached results in the serialized workbook. openpyxl does not calculate them.
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cached = [("str", "N/A"), ("n", "7"), ("e", "#N/A"), ("str", None)]
    result = io.BytesIO()
    with ZipFile(io.BytesIO(stream.getvalue())) as source, ZipFile(result, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                for index, (kind, value) in enumerate(cached, 1):
                    cell = root.find(f'.//{namespace}c[@r="B{index}"]')
                    cell.set("t", kind)
                    cell.find(f"{namespace}v").text = value
                content = ElementTree.tostring(root)
            target.writestr(item, content)
    with open_excel_workbook(result.getvalue()) as reader:
        frame = reader.parse("Sheet")
    assert frame.iloc[0, 1] == "N/A"
    assert frame.iloc[1, 1] == 7
    assert pd.isna(frame.iloc[2, 1])
    assert pd.isna(frame.iloc[3, 1])


def test_shared_xls_reader_preserves_tokens_and_blank_separator():
    # Serialized with xlwt: text tokens in B2/B3/B5/B6 and a physically blank fourth row.
    workbook = zlib.decompress(
        base64.b64decode(
            "eJztWE1oE0EU/mazm5/SpmlNhVYooWDVWhDx4qVdI9giaog/oIigqQ0SkiaSJoherNYcBcGT4qXQi5eqF3/Qg948"
            "CBU9CIKQ6tFTQcFDm/W9t7uYaoTm0PrDfmHezLz5Zt7LzJvZ2X0931GdedCzgJ8wDB9qVgj+Op2iFHIrEVC7ZXHR"
            "zYOULA//FEJBWki/gadtrwK8hrzeC9BwX39BEvhI6RTOI1HIp2PriL3iQ0qxD0MkFe6QJoxu8apT5FmRG0TeE+Yz"
            "kXtEc13kEHGr6iTmzcTAbieKT2h90hYGj/tI+rwXzU504SVH8eUbyuYaiBczqdzf2dCrt2IWtG6j6Xy6mMpVEaUF"
            "nMVXKwZ8cXfq85inX1+9Aum/rdQHGuhvajowBeuMBHiFAnLRxy1+HCtk0/nJJRxEC9U5EXN/KT1B+/NoKVUqT1I9"
            "mU4ViZ/YESc5kjlHEZyIsz5XnqCMtyxlBzIXMpTly7kcj2Nv+MiKDd8mG6GV5Djapdwh2yFCR/rS3cU3h8aS5mnR"
            "TMkhbz8KNrP3sHCFe1DnMNzxINwBkdtFXpVRN0m5R2SUgpjy/mSXUxiZFs41ae0nO7sEb80tdeWtVK58Pvy4t/LJ"
            "3EbludGFS9G5d+YM+ujRNE79+TeNQTWobt9iPDHdXDnHxgeR3b8cIUEt4vhu2Quk2rEscw+aD4Zd49lRDl814Cvh"
            "aw5fCd/n8LUGfE34usO3Z99w+HoDvi58v8PXhR9w+EYDviH8oMM3hM+P8WGtEw/ZOB2YP9ACDx48ePDgwYMHD/8x"
            "7Ass5ILKF0u+LPLVki+U/F1nmVLtT3+k8LBmOIIC/Ur0YroPecqLuNhU/GyEodyx1Cr7uN8LGcfJehFZjIkf2aZs"
            "M+iVTtX/n1V3jDRt6ndo2n6tGT/X2P53eZHVvg=="
        )
    )
    with open_excel_workbook(workbook) as reader:
        frame = reader.parse("Tokens")
    assert frame.iloc[[1, 2, 4, 5], 1].tolist() == ["N/A", "NA", "None", "null"]
    assert frame.iloc[3].isna().all()
    assert [element.text for element in partition_xlsx(file=io.BytesIO(workbook))] == [
        "Item Status Pear N/A Fig NA",
        "Plum None Kiwi null",
    ]
