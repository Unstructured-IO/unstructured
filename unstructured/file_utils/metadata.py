import datetime
import io
from dataclasses import dataclass, field
from typing import IO, Any, Dict, Final, Iterable, Iterator, List, Optional, Union

import docx
import openpyxl
from PIL import Image
from PIL.ExifTags import TAGS

from unstructured.documents.elements import Element
from unstructured.partition.lang import detect_languages

# NOTE(robison) - ref: https://www.media.mit.edu/pia/Research/deepview/exif.html
EXIF_DATETIME_FMT: Final[str] = "%Y:%m:%d %H:%M:%S"


@dataclass
class Metadata:
    author: str = ""
    category: str = ""
    comments: str = ""
    content_status: str = ""
    created: Optional[datetime.datetime] = None
    identifier: str = ""
    keywords: str = ""
    language: str = ""
    last_modified_by: str = ""
    last_printed: Optional[datetime.datetime] = None
    modified: Optional[datetime.datetime] = None
    revision: Optional[int] = 0
    subject: str = ""
    title: str = ""
    version: str = ""
    description: str = ""
    namespace: str = ""

    # NOTE(robinson) - Metadata for use with image files
    exif_data: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self):
        return self.__dict__


def get_docx_metadata(
    filename: str = "",
    file: Optional[IO[bytes]] = None,
) -> Metadata:
    """Extracts document metadata from a Microsoft .docx document."""
    if filename:
        doc = docx.Document(filename)
    elif file:
        doc = docx.Document(file)
    else:
        raise FileNotFoundError("No filename nor file were specified")

    metadata = Metadata(
        author=getattr(doc.core_properties, "author", ""),
        category=getattr(doc.core_properties, "category", ""),
        comments=getattr(doc.core_properties, "comments", ""),
        content_status=getattr(doc.core_properties, "content_status", ""),
        created=getattr(doc.core_properties, "created", None),
        identifier=getattr(doc.core_properties, "identifier", ""),
        keywords=getattr(doc.core_properties, "keywords", ""),
        language=getattr(doc.core_properties, "language", ""),
        last_modified_by=getattr(doc.core_properties, "last_modified_by", ""),
        last_printed=getattr(doc.core_properties, "last_printed", None),
        modified=getattr(doc.core_properties, "modified", None),
        revision=getattr(doc.core_properties, "revision", None),
        subject=getattr(doc.core_properties, "subject", ""),
        title=getattr(doc.core_properties, "title", ""),
        version=getattr(doc.core_properties, "version", ""),
    )

    return metadata


def get_xlsx_metadata(
    filename: str = "",
    file: Optional[IO[bytes]] = None,
) -> Metadata:
    """Extracts document metadata from a Microsoft .xlsx document."""
    if filename:
        workbook = openpyxl.load_workbook(filename)
    elif file:
        workbook = openpyxl.load_workbook(file)
    else:
        raise FileNotFoundError("No filename nor file were specified")

    metadata = Metadata(
        author=getattr(workbook.properties, "creator", ""),
        category=getattr(workbook.properties, "category", ""),
        content_status=getattr(workbook.properties, "contentStatus", ""),
        created=getattr(workbook.properties, "created", None),
        description=getattr(workbook.properties, "description", ""),
        identifier=getattr(workbook.properties, "identifier", ""),
        keywords=getattr(workbook.properties, "keywords", ""),
        language=getattr(workbook.properties, "language", ""),
        last_modified_by=getattr(workbook.properties, "lastModifiedBy", ""),
        last_printed=getattr(workbook.properties, "lastPrinted", None),
        modified=getattr(workbook.properties, "modified", None),
        namespace=getattr(workbook.properties, "namespace", ""),
        revision=getattr(workbook.properties, "revision", None),
        subject=getattr(workbook.properties, "subject", ""),
        title=getattr(workbook.properties, "title", ""),
        version=getattr(workbook.properties, "version", ""),
    )

    return metadata


def get_jpg_metadata(
    filename: str = "",
    file: Optional[IO[bytes]] = None,
) -> Metadata:
    """Extracts metadata from a JPG image, including EXIF metadata."""
    if filename:
        image = Image.open(filename)
    elif file:
        image = Image.open(io.BytesIO(file.read()))
    else:
        raise FileNotFoundError("No filename nor file were specified")

    exif_data = image.getexif()
    exif_dict: Dict[str, Any] = {}
    for tag_id in exif_data:
        tag = TAGS.get(tag_id, tag_id)
        data = exif_data.get(tag_id)
        exif_dict[tag] = data

    metadata = Metadata(
        author=exif_dict.get("Artist", ""),
        comments=exif_dict.get("UserComment", ""),
        created=_get_exif_datetime(exif_dict, "DateTimeOriginal"),
        # NOTE(robinson) - Per EXIF docs, DateTime is the last modified data
        # ref: https://www.media.mit.edu/pia/Research/deepview/exif.html
        modified=_get_exif_datetime(exif_dict, "DateTime"),
        exif_data=exif_dict,
    )

    return metadata


def _get_exif_datetime(exif_dict: Dict[str, Any], key: str) -> Optional[datetime.datetime]:
    """Converts a datetime string from the EXIF data to a Python datetime object."""
    date = exif_dict.get(key)
    if not date:
        return None

    try:
        return datetime.datetime.strptime(date, EXIF_DATETIME_FMT)
    # NOTE(robinson) - An exception could occur if the datetime is not formatted
    # using the standard EXIF datetime format
    except ValueError:
        return None


def apply_lang_metadata(
    elements: Union[Iterable[Element], List[Element]],
    languages: List[str],
    detect_language_per_element: bool = False,
) -> Iterator[Element]:
    """Detect and apply metadata.languages to each element in `elements`."""
    # -- Note this function has a stream interface, but reads the full `elements` stream into memory
    # -- before emitting the first updated element as output.

    # Skip language detection for partitioners that use other partitioners.
    # For example, partition_msg relies on partition_html and partition_text, but the metadata
    # gets overwritten after elements have been returned by _html and _text,
    # so `languages` would be detected twice.
    if languages == [""]:
        yield from elements
        return

    if not isinstance(elements, List):
        elements = list(elements)

    full_text = " ".join(e.text for e in elements if hasattr(e, "text"))
    detected_languages = detect_languages(text=full_text, languages=languages)
    if (
        detected_languages is not None
        and len(languages) == 1
        and detect_language_per_element is False
    ):
        # -- apply detected languge to each metadata --
        for e in elements:
            e.metadata.languages = detected_languages
            yield e
    else:
        for e in elements:
            if not hasattr(e, "text"):
                yield e
                continue
            e.metadata.languages = detect_languages(e.text)
            yield e
